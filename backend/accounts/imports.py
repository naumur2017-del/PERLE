"""Import en masse par fichier Excel pour les pages Architecture (catalogue des tâches et
architecture monétaire) : chaque ligne du classeur est validée et enregistrée via le même
serializer qu'une création manuelle (TaskTemplateSerializer / LigneBudgetaireCreateSerializer),
pour ne jamais dupliquer les règles métier. Les lignes sont résolues en plusieurs passes afin
d'accepter un classeur où un parent apparaît après ses enfants."""
import io
import re
import unicodedata
from decimal import Decimal, InvalidOperation

import openpyxl
from openpyxl.styles import Font
from rest_framework.exceptions import PermissionDenied, ValidationError
from rest_framework.parsers import FormParser, MultiPartParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response
from rest_framework.views import APIView

from .access import can_access_config
from .models import TASK_PRIORITE_CHOICES, LigneBudgetaire, TaskTemplate, Team
from .serializers import LigneBudgetaireCreateSerializer, LigneBudgetaireSerializer, TaskTemplateSerializer


def _norm(text):
    text = unicodedata.normalize('NFKD', str(text or ''))
    text = ''.join(char for char in text if not unicodedata.combining(char)).lower()
    return re.sub(r'[^a-z0-9]+', ' ', text).strip()


def _read_rows(fichier):
    if not fichier.name.lower().endswith(('.xlsx', '.xlsm')):
        raise ValidationError('Le fichier doit être un classeur Excel (.xlsx).')
    try:
        classeur = openpyxl.load_workbook(fichier, data_only=True, read_only=True)
    except Exception:
        raise ValidationError('Fichier Excel illisible ou corrompu.')
    feuille = classeur.active
    toutes_lignes = list(feuille.iter_rows(values_only=True))
    classeur.close()
    if not toutes_lignes:
        raise ValidationError('Le fichier est vide.')
    entetes = [_norm(cell) for cell in toutes_lignes[0]]
    rows = []
    for numero, row in enumerate(toutes_lignes[1:], start=2):
        if all(cell is None or str(cell).strip() == '' for cell in row):
            continue
        rows.append((numero, {entetes[i]: row[i] for i in range(min(len(entetes), len(row)))}))
    return rows


def _cell(valeurs, *alias):
    for cle in alias:
        if cle in valeurs and valeurs[cle] is not None:
            texte = str(valeurs[cle]).strip()
            if texte:
                return texte
    return ''


def _bool_cell(valeurs, default, *alias):
    texte = _cell(valeurs, *alias).lower()
    if not texte:
        return default
    return texte in ('oui', 'o', 'yes', 'y', 'true', '1', 'x')


def _decimal_cell(valeurs, *alias):
    texte = _cell(valeurs, *alias)
    if not texte:
        return None
    try:
        return Decimal(texte.replace(',', '.').replace(' ', ''))
    except InvalidOperation:
        raise ValueError(f'« {texte} » n’est pas un nombre valide.')


def _choice_cell(choices, valeurs, default, *alias):
    texte = _cell(valeurs, *alias)
    if not texte:
        return default
    texte_norm = _norm(texte)
    for key, label in choices:
        if _norm(key) == texte_norm or _norm(label) == texte_norm:
            return key
    valeurs_possibles = ', '.join(label for _key, label in choices)
    raise ValueError(f'« {texte} » n’est pas reconnu (valeurs possibles : {valeurs_possibles}).')


def _resoudre_equipe(organisation, texte):
    if not texte:
        return None
    return (
        Team.objects.filter(organisation=organisation, code__iexact=texte).first()
        or Team.objects.filter(organisation=organisation, name__iexact=texte).first()
    )


def _resolution_par_passes(en_attente, resoudre_ligne):
    """Traite `en_attente` (code -> (numero, valeurs)) en plusieurs passes, pour que l'ordre des
    lignes dans le classeur n'ait pas d'importance : une ligne n'est traitée que lorsque son code
    parent a déjà été résolu (préexistant en base ou créé dans une passe précédente).
    `resoudre_ligne(code, numero, valeurs)` doit renvoyer True si la ligne a été traitée
    (créée ou en erreur), False si elle doit attendre une prochaine passe."""
    progresse = True
    while en_attente and progresse:
        progresse = False
        for code in list(en_attente.keys()):
            numero, valeurs = en_attente[code]
            if resoudre_ligne(code, numero, valeurs):
                del en_attente[code]
                progresse = True
    return en_attente


class TaskTemplateImportView(APIView):
    """Import du catalogue des tâches (Architecture des tâches). GET : télécharge un modèle
    Excel vierge. POST (multipart, champ `file`) : importe le classeur."""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    HEADERS = ['Code', 'Nom', 'Code parent', 'Equipe', 'Type (Dossier / Tâche élémentaire)',
               'Attribuable (Oui/Non)', 'Récurrente (Oui/Non)', 'Fréquence (Ponctuelle/Récurrente)',
               'Déclenchement (Manuel/Automatique)', 'Priorité (Haute/Moyenne/Basse)',
               'Durée estimée (h)', 'Détails', 'Explication']

    def get(self, request):
        if not can_access_config(request.user):
            raise PermissionDenied('Vous n’êtes pas autorisé à importer le catalogue des tâches.')
        exemples = [
            ['PIL', 'Pilotage', '', 'PIL', 'Dossier', 'Oui', 'Non', 'Ponctuelle', 'Manuel', 'Moyenne', '', 'Dossier racine de l’équipe Pilotage', ''],
            ['PIL-01', 'Chiffrage unitaire', 'PIL', '', 'Tâche élémentaire', 'Oui', 'Non', 'Ponctuelle', 'Manuel', 'Moyenne', '2', 'Exemple de tâche élémentaire', ''],
        ]
        return _modele_xlsx('catalogue-des-taches-modele.xlsx', self.HEADERS, exemples)

    def post(self, request):
        if not can_access_config(request.user):
            raise PermissionDenied('Vous n’êtes pas autorisé à importer le catalogue des tâches.')
        organisation = request.user.organisation
        if not organisation:
            raise PermissionDenied('Votre compte n’est rattaché à aucune organisation.')
        fichier = request.FILES.get('file')
        if not fichier:
            raise ValidationError('Aucun fichier n’a été envoyé.')

        codes_connus = {c.lower() for c in TaskTemplate.objects.filter(organisation=organisation).values_list('code', flat=True)}
        en_attente = {}
        codes_vus = set()
        errors = []
        for numero, valeurs in _read_rows(fichier):
            code = _cell(valeurs, 'code')
            if not code:
                errors.append({'ligne': numero, 'code': '', 'erreurs': {'code': 'Le code est obligatoire.'}})
                continue
            if code.lower() in codes_vus:
                errors.append({'ligne': numero, 'code': code, 'erreurs': {'code': 'Ce code apparaît plusieurs fois dans le fichier.'}})
                continue
            codes_vus.add(code.lower())
            en_attente[code] = (numero, valeurs)

        created = []

        def resoudre(code, numero, valeurs):
            code_parent = _cell(valeurs, 'code parent', 'parent')
            if code_parent and code_parent.lower() not in codes_connus:
                return False
            try:
                data = {
                    'code': code,
                    'nom': _cell(valeurs, 'nom'),
                    'parent': None,
                    'equipe': None,
                    'type_element': _choice_cell(TaskTemplate.TYPE_CHOICES, valeurs, 'dossier', 'type', 'type dossier tache elementaire'),
                    'attribuable': _bool_cell(valeurs, True, 'attribuable', 'attribuable oui non'),
                    'recurrente': _bool_cell(valeurs, False, 'recurrente', 'recurrente oui non'),
                    'frequence': _choice_cell(TaskTemplate.FREQUENCE_CHOICES, valeurs, 'ponctuelle', 'frequence', 'frequence ponctuelle recurrente'),
                    'mode_declenchement': _choice_cell(TaskTemplate.DECLENCHEMENT_CHOICES, valeurs, 'manuel', 'declenchement', 'mode declenchement', 'declenchement manuel automatique'),
                    'priorite_defaut': _choice_cell(TASK_PRIORITE_CHOICES, valeurs, 'moyenne', 'priorite', 'priorite defaut', 'priorite haute moyenne basse'),
                    'duree_estimee_heures': _decimal_cell(valeurs, 'duree estimee h', 'duree estimee', 'duree'),
                    'details': _cell(valeurs, 'details'),
                    'explication': _cell(valeurs, 'explication'),
                }
            except ValueError as exc:
                errors.append({'ligne': numero, 'code': code, 'erreurs': {'detail': str(exc)}})
                return True

            if code_parent:
                parent = TaskTemplate.objects.filter(organisation=organisation, code__iexact=code_parent).first()
                data['parent'] = parent.id if parent else None
            equipe = _resoudre_equipe(organisation, _cell(valeurs, 'equipe'))
            if equipe is not None:
                data['equipe'] = equipe.id

            serializer = TaskTemplateSerializer(data=data, context={'request': request})
            if serializer.is_valid():
                serializer.save()
                created.append(serializer.instance)
                codes_connus.add(code.lower())
            else:
                errors.append({'ligne': numero, 'code': code, 'erreurs': serializer.errors})
            return True

        restants = _resolution_par_passes(en_attente, resoudre)
        for code, (numero, _valeurs) in restants.items():
            errors.append({'ligne': numero, 'code': code, 'erreurs': {'parent': 'Code parent introuvable (dépendance manquante ou circulaire).'}})

        return Response({
            'created': len(created),
            'errors': errors,
            'items': TaskTemplateSerializer(created, many=True, context={'request': request}).data,
        }, status=201 if created else 200)


class LigneBudgetaireImportView(APIView):
    """Import de l'architecture monétaire (lignes budgétaires). GET : télécharge un modèle
    Excel vierge. POST (multipart, champ `file`) : importe le classeur."""
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    HEADERS = ['Code', 'Nom', 'Code parent', 'Equipe', 'Déclinaison', 'Montant prévu']

    def get(self, request):
        if not can_access_config(request.user):
            raise PermissionDenied('Vous n’êtes pas autorisé à importer l’architecture monétaire.')
        exemples = [
            ['A', 'Fonctionnement', '', 'RES', '', ''],
            ['AA', 'Carburant', 'A', '', 'Parc automobile', '1000000'],
        ]
        return _modele_xlsx('architecture-monetaire-modele.xlsx', self.HEADERS, exemples)

    def post(self, request):
        if not can_access_config(request.user):
            raise PermissionDenied('Vous n’êtes pas autorisé à importer l’architecture monétaire.')
        organisation = request.user.organisation
        if not organisation:
            raise PermissionDenied('Votre compte n’est rattaché à aucune organisation.')
        fichier = request.FILES.get('file')
        if not fichier:
            raise ValidationError('Aucun fichier n’a été envoyé.')

        codes_connus = {c.lower() for c in LigneBudgetaire.objects.filter(organisation=organisation).values_list('code', flat=True)}
        equipe_id_par_code = dict(
            (code.lower(), equipe_id)
            for code, equipe_id in LigneBudgetaire.objects.filter(organisation=organisation).values_list('code', 'equipe_id')
        )
        en_attente = {}
        codes_vus = set()
        errors = []
        for numero, valeurs in _read_rows(fichier):
            code = _cell(valeurs, 'code')
            if not code:
                errors.append({'ligne': numero, 'code': '', 'erreurs': {'code': 'Le code est obligatoire.'}})
                continue
            if code.lower() in codes_vus:
                errors.append({'ligne': numero, 'code': code, 'erreurs': {'code': 'Ce code apparaît plusieurs fois dans le fichier.'}})
                continue
            codes_vus.add(code.lower())
            en_attente[code] = (numero, valeurs)

        created = []

        def resoudre(code, numero, valeurs):
            code_parent = _cell(valeurs, 'code parent', 'parent')
            if code_parent and code_parent.lower() not in codes_connus:
                return False
            try:
                data = {
                    'code': code,
                    'nom': _cell(valeurs, 'nom'),
                    'parent': None,
                    'equipe': None,
                    'declinaison': _cell(valeurs, 'declinaison'),
                    'montant_prevu': _decimal_cell(valeurs, 'montant prevu', 'montant'),
                }
            except ValueError as exc:
                errors.append({'ligne': numero, 'code': code, 'erreurs': {'detail': str(exc)}})
                return True

            if code_parent:
                parent = LigneBudgetaire.objects.filter(organisation=organisation, code__iexact=code_parent).first()
                data['parent'] = parent.id if parent else None
            equipe = _resoudre_equipe(organisation, _cell(valeurs, 'equipe'))
            if equipe is None and code_parent:
                # Équipe non précisée sur une sous-ligne : hérite de la ligne parente, comme le
                # catalogue des tâches (voir TaskTemplateSerializer.create).
                parent_equipe_id = equipe_id_par_code.get(code_parent.lower())
                if parent_equipe_id:
                    equipe = Team.objects.filter(pk=parent_equipe_id).first()
            if equipe is not None:
                data['equipe'] = equipe.id

            serializer = LigneBudgetaireCreateSerializer(data=data, context={'request': request})
            if serializer.is_valid():
                serializer.save()
                created.append(serializer.instance)
                codes_connus.add(code.lower())
                equipe_id_par_code[code.lower()] = serializer.instance.equipe_id
            else:
                errors.append({'ligne': numero, 'code': code, 'erreurs': serializer.errors})
            return True

        restants = _resolution_par_passes(en_attente, resoudre)
        for code, (numero, _valeurs) in restants.items():
            errors.append({'ligne': numero, 'code': code, 'erreurs': {'parent': 'Code parent introuvable (dépendance manquante ou circulaire).'}})

        return Response({
            'created': len(created),
            'errors': errors,
            'items': LigneBudgetaireSerializer(created, many=True, context={'request': request}).data,
        }, status=201 if created else 200)


def _modele_xlsx(nom_fichier, headers, exemples):
    from django.http import HttpResponse

    classeur = openpyxl.Workbook()
    feuille = classeur.active
    feuille.append(headers)
    for cell in feuille[1]:
        cell.font = Font(bold=True)
    for row in exemples:
        feuille.append(row)
    for colonne in feuille.columns:
        largeur = max(len(str(cell.value)) if cell.value is not None else 0 for cell in colonne)
        feuille.column_dimensions[colonne[0].column_letter].width = min(max(largeur + 2, 12), 40)

    tampon = io.BytesIO()
    classeur.save(tampon)
    tampon.seek(0)
    response = HttpResponse(
        tampon.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    )
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'
    return response
