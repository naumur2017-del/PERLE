"""Exports tabulaires communs à toutes les pages (CSV, Excel ou PDF).

Le frontend envoie le titre, la période, les KPI et les lignes telles qu'elles sont affichées
(déjà formatées), et reçoit un fichier prêt à imprimer. Mise en page commune : A4 paysage, colonnes
ajustées à la largeur de la page, texte qui passe à la ligne, entête répété sur chaque page."""

import csv
import io
import re
import unicodedata
from decimal import Decimal

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from rest_framework import serializers
from rest_framework.permissions import IsAuthenticated
from rest_framework.views import APIView

MIME = {
    'csv': 'text/csv; charset=utf-8',
    'xlsx': 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
    'pdf': 'application/pdf',
}
NUMERIQUE = re.compile(r'^[-+]?[\d\s.,]+%?$')
VIOLET = '6B46C1'


def _est_numerique(texte):
    return bool(NUMERIQUE.match(texte.strip())) and any(c.isdigit() for c in texte)


def _nombre(texte):
    """« 1 234,50 » -> 1234.5 (None si ce n'est pas un nombre)."""
    propre = texte.replace(' ', '').replace(' ', '').replace('%', '').replace(',', '.')
    try:
        return float(propre)
    except ValueError:
        return None


def _nom_fichier(nom):
    ascii_ = unicodedata.normalize('NFKD', nom).encode('ascii', 'ignore').decode()
    return re.sub(r'[^A-Za-z0-9_-]+', '-', ascii_).strip('-') or 'export'


def _colonnes_numeriques(colonnes, lignes):
    """Index des colonnes à aligner à droite : titre « montant/total/solde/nombre » ou valeurs numériques."""
    indices = set()
    for index, titre in enumerate(colonnes):
        if re.search(r'montant|total|solde|nombre', titre, re.IGNORECASE):
            indices.add(index)
            continue
        valeurs = [ligne[index] for ligne in lignes if index < len(ligne) and ligne[index].strip() not in ('', '—')]
        if valeurs and all(_est_numerique(v) for v in valeurs):
            indices.add(index)
    return indices


def render_csv(titre, periode, kpis, colonnes, lignes):
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=';')
    writer.writerow([titre])
    if periode:
        writer.writerow([periode])
    for label, valeur in kpis:
        writer.writerow([label, valeur])
    writer.writerow([])
    writer.writerow(colonnes)
    writer.writerows(lignes)
    # BOM UTF-8 : Excel ouvre directement le fichier avec les accents.
    return buffer.getvalue().encode('utf-8-sig')


def render_xlsx(titre, periode, kpis, colonnes, lignes):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Export'
    ws['A1'] = titre
    ws['A1'].font = Font(bold=True, size=14)
    row = 2
    if periode:
        ws.cell(row=row, column=1, value=periode).font = Font(italic=True)
        row += 1
    row += 1
    for label, valeur in kpis:
        ws.cell(row=row, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row, column=2, value=valeur)
        row += 1
    row += 1
    entete_row = row
    for col, titre_col in enumerate(colonnes, start=1):
        cell = ws.cell(row=row, column=col, value=titre_col)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor=VIOLET)
        cell.alignment = Alignment(horizontal='center', wrap_text=True)
    numeriques = _colonnes_numeriques(colonnes, lignes)
    for ligne in lignes:
        row += 1
        for col, valeur in enumerate(ligne, start=1):
            nombre = _nombre(valeur) if (col - 1) in numeriques else None
            cell = ws.cell(row=row, column=col, value=nombre if nombre is not None else valeur)
            if nombre is not None:
                cell.number_format = '#,##0.00'
            cell.alignment = Alignment(horizontal='right' if col - 1 in numeriques else 'left', vertical='top', wrap_text=True)
    for col, titre_col in enumerate(colonnes, start=1):
        longueur = max([len(titre_col)] + [len(ligne[col - 1]) for ligne in lignes if col - 1 < len(ligne)])
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = min(max(10, longueur + 2), 45)
    ws.freeze_panes = ws.cell(row=entete_row + 1, column=1)
    # Impression : A4 paysage, largeur sur une page, entête répété.
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = f'{entete_row}:{entete_row}'
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def render_pdf(titre, periode, kpis, colonnes, lignes):
    out = io.BytesIO()
    marge = 24
    doc = SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=marge, rightMargin=marge, topMargin=marge, bottomMargin=marge)
    largeur_utile = landscape(A4)[0] - 2 * marge
    styles = getSampleStyleSheet()
    violet = colors.HexColor('#' + VIOLET)
    elements = [Paragraph(titre, styles['Title'])]
    if periode:
        elements.append(Paragraph(periode, styles['Normal']))
    elements.append(Spacer(1, 8))

    if kpis:
        kpi_table = Table([[label, valeur] for label, valeur in kpis], colWidths=[220, 160], hAlign='LEFT')
        kpi_table.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 10),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('LINEBELOW', (0, 0), (-1, -1), 0.25, colors.HexColor('#EAE7F6')),
        ]))
        elements += [kpi_table, Spacer(1, 14)]

    # Largeur de chaque colonne proportionnelle à son contenu le plus long (bornée), puis ajustée à la page.
    poids = []
    for index, titre_col in enumerate(colonnes):
        longueur = max([len(titre_col)] + [len(ligne[index]) for ligne in lignes if index < len(ligne)])
        poids.append(min(max(longueur, 6), 45))
    facteur = largeur_utile / sum(poids)
    largeurs = [p * facteur for p in poids]

    numeriques = _colonnes_numeriques(colonnes, lignes)
    cellule = ParagraphStyle('cellule', fontName='Helvetica', fontSize=7.5, leading=9)
    cellule_num = ParagraphStyle('num', parent=cellule, alignment=TA_RIGHT)
    cellule_entete = ParagraphStyle('entete', parent=cellule, fontName='Helvetica-Bold', textColor=colors.white)
    donnees = [[Paragraph(titre_col, cellule_entete) for titre_col in colonnes]]
    for ligne in lignes:
        donnees.append([
            Paragraph(valeur, cellule_num if index in numeriques else cellule)
            for index, valeur in enumerate(ligne)
        ])
    table = Table(donnees, colWidths=largeurs, repeatRows=1)
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), violet),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#FAF9FF')]),
        ('GRID', (0, 0), (-1, -1), 0.25, colors.HexColor('#EAE7F6')),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('LEFTPADDING', (0, 0), (-1, -1), 3),
        ('RIGHTPADDING', (0, 0), (-1, -1), 3),
    ]))
    elements.append(table)
    doc.build(elements)
    return out.getvalue()


RENDERERS = {'csv': render_csv, 'xlsx': render_xlsx, 'pdf': render_pdf}


def reponse_tableau(format_, nom, titre, periode, kpis, colonnes, lignes):
    """Fichier téléchargeable. `kpis` : liste de (libellé, valeur) ; `lignes` : valeurs déjà formatées."""
    contenu = RENDERERS[format_](titre, periode, kpis, colonnes, lignes)
    nom_fichier = f'{_nom_fichier(nom)}.{format_}'
    response = HttpResponse(contenu, content_type=MIME[format_])
    response['Content-Disposition'] = f'attachment; filename="{nom_fichier}"'
    return response


class TableauExportSerializer(serializers.Serializer):
    format = serializers.ChoiceField(choices=list(MIME))
    nom = serializers.CharField(max_length=120)
    titre = serializers.CharField(max_length=200)
    periode = serializers.CharField(max_length=200, required=False, allow_blank=True, default='')
    kpis = serializers.ListField(
        child=serializers.ListField(child=serializers.CharField(max_length=200, allow_blank=True), min_length=2, max_length=2),
        max_length=30, required=False, default=list,
    )
    colonnes = serializers.ListField(child=serializers.CharField(max_length=100), min_length=1, max_length=20)
    lignes = serializers.ListField(
        child=serializers.ListField(child=serializers.CharField(max_length=500, allow_blank=True), max_length=20),
        max_length=10000, allow_empty=True,
    )

    def validate(self, attrs):
        colonnes = attrs['colonnes']
        for ligne in attrs['lignes']:
            if len(ligne) != len(colonnes):
                raise serializers.ValidationError({'lignes': 'Chaque ligne doit avoir autant de valeurs que de colonnes.'})
        return attrs


class TableauExportView(APIView):
    """POST : télécharge le tableau affiché (CSV, Excel ou PDF). Ouvert à tout utilisateur connecté :
    les données envoyées sont celles que la page lui affiche déjà."""
    permission_classes = [IsAuthenticated]

    def post(self, request):
        serializer = TableauExportSerializer(data=request.data)
        serializer.is_valid(raise_exception=True)
        d = serializer.validated_data
        return reponse_tableau(
            d['format'], d['nom'], d['titre'], d['periode'],
            [(k[0], k[1]) for k in d['kpis']], d['colonnes'], d['lignes'],
        )
