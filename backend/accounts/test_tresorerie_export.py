import csv
import io

from openpyxl import load_workbook
from rest_framework.test import APITestCase

from .models import CompteTresorerie, MouvementTresorerie, Organisation, User


class JournalExportTests(APITestCase):
    def setUp(self):
        self.org = Organisation.objects.create(name='Export SA', org_type='company', currency_code='XAF')
        self.other = Organisation.objects.create(name='Autre', org_type='company')
        self.director = User.objects.create_user(email='dir@export.test', password='x', role='directeur', organisation=self.org)
        self.employee = User.objects.create_user(email='emp@export.test', password='x', role='salarie', organisation=self.org)
        self.compte = CompteTresorerie.objects.create(organisation=self.org, nom='BICEC', code='521100', solde_initial=0)
        self.entree = MouvementTresorerie.objects.create(organisation=self.org, compte=self.compte, type_mouvement='Entrée',
                                                         nature='Approvisionnement', libelle='Apport', montant=1000, origine='Siège')
        self.sortie = MouvementTresorerie.objects.create(organisation=self.org, compte=self.compte, type_mouvement='Sortie',
                                                         nature='Paiement', libelle='Achat', montant=250, beneficiaire='ETS X')
        autre = CompteTresorerie.objects.create(organisation=self.other, nom='Autre', code='1', solde_initial=0)
        self.intrus = MouvementTresorerie.objects.create(organisation=self.other, compte=autre, type_mouvement='Entrée',
                                                         nature='Approvisionnement', libelle='Intrus', montant=9999)
        self.client.force_authenticate(self.director)

    def export(self, format_, ids=None, **extra):
        body = {'format': format_, 'date_debut': '2026-01-01', 'date_fin': '2026-12-31',
                'mouvement_ids': ids if ids is not None else [self.entree.pk, self.sortie.pk, self.intrus.pk], **extra}
        return self.client.post('/api/tresorerie/journal/export/', body, format='json')

    def test_csv_contains_kpis_and_only_organisation_rows(self):
        response = self.export('csv')
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment; filename="journal-tresorerie_2026-01-01_2026-12-31.csv"', response['Content-Disposition'])
        texte = response.content.decode('utf-8-sig')
        self.assertIn('Total entrées;1 000,00', texte)
        self.assertIn('Total sorties;250,00', texte)
        self.assertIn('Solde net de la période;750,00', texte)
        self.assertIn("Nombre d'opérations;2", texte)
        self.assertIn('Apport', texte)
        self.assertNotIn('Intrus', texte)
        lignes = list(csv.reader(io.StringIO(texte), delimiter=';'))
        entete = lignes[lignes.index(['Date', 'Heure', 'Référence', 'Code projet', 'Libellé', 'Type', 'Compte',
                                      'Ordonnateur', 'Bénéficiaire', 'Montant', 'Justificatif', 'Exécuteur'])]
        self.assertEqual(entete[0], 'Date')

    def test_xlsx_is_a_workbook_with_kpis_and_rows(self):
        response = self.export('xlsx')
        self.assertEqual(response.status_code, 200)
        feuille = load_workbook(io.BytesIO(response.content)).active
        valeurs = [cell.value for row in feuille.iter_rows() for cell in row]
        self.assertIn('Total entrées', valeurs)
        self.assertIn(1000.0, valeurs)
        self.assertIn('Apport', valeurs)
        self.assertNotIn('Intrus', valeurs)

    def test_pdf_is_produced(self):
        response = self.export('pdf')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response['Content-Type'], 'application/pdf')
        self.assertTrue(response.content.startswith(b'%PDF'))

    def test_only_selected_rows_are_exported(self):
        texte = self.export('csv', ids=[self.sortie.pk]).content.decode('utf-8-sig')
        self.assertIn("Nombre d'opérations;1", texte)
        self.assertNotIn('Apport', texte)

    def test_invalid_requests_and_permissions(self):
        self.assertEqual(self.export('docx').status_code, 400)
        self.assertEqual(self.export('csv', date_fin='2025-01-01').status_code, 400)
        self.client.force_authenticate(self.employee)
        self.assertEqual(self.export('csv').status_code, 403)

    def test_pdf_is_a4_landscape(self):
        import re
        contenu = self.export('pdf').content
        largeur, hauteur = map(float, re.search(rb'/MediaBox\s*\[\s*0\s+0\s+([\d.]+)\s+([\d.]+)\s*\]', contenu).groups())
        self.assertAlmostEqual(largeur, 841.89, places=1)
        self.assertAlmostEqual(hauteur, 595.28, places=1)

    def test_xlsx_is_set_for_a4_landscape_print(self):
        feuille = load_workbook(io.BytesIO(self.export('xlsx').content)).active
        self.assertEqual(str(feuille.page_setup.paperSize), str(feuille.PAPERSIZE_A4))
        self.assertEqual(feuille.page_setup.orientation, feuille.ORIENTATION_LANDSCAPE)
        self.assertEqual(feuille.page_setup.fitToWidth, 1)
